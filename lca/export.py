"""Export an assessment to XLSX — the format the SGC IS tender itself requires for list exports."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .model import MODULES, Building, assess

_HEAD = Font(bold=True, color="FFFFFF")
_HFILL = PatternFill("solid", fgColor="1D4ED8")
_TITLE = Font(bold=True, size=14, color="1D4ED8")


def to_xlsx(b: Building) -> bytes:
    r = assess(b)
    wb = Workbook()

    ws = wb.active
    ws.title = "RESULTS"
    ws["A1"] = "FastLCA — Whole-building life-cycle carbon (EN 15978)"
    ws["A1"].font = _TITLE
    meta = [("Building type", b.building_type), ("Gross floor area (m²)", b.gross_floor_area),
            ("Year of completion", b.year_completion), ("Assessment period (years)", b.assessment_period)]
    for i, (k, v) in enumerate(meta, start=3):
        ws[f"A{i}"] = k; ws[f"A{i}"].font = Font(bold=True); ws[f"B{i}"] = v

    row = 8
    for col, h in enumerate(["Module", "Description", "tCO₂e", "Share of total (%)"], start=1):
        c = ws.cell(row, col, h); c.font = _HEAD; c.fill = _HFILL
    row += 1
    total = r["totals"]["total_excl_d"] or 1
    labels = dict(MODULES)
    for code, val in r["modules"].items():
        ws.cell(row, 1, code)
        ws.cell(row, 2, labels.get(code, ""))
        ws.cell(row, 3, val)
        ws.cell(row, 4, round(val / total * 100, 1) if code != "D1" else None)
        row += 1
    row += 1
    for k, label in [("embodied", "Embodied carbon (A+C, excl. B6/B7)"),
                     ("operational", "Operational carbon (B6+B7)"),
                     ("total_excl_d", "TOTAL excl. module D"),
                     ("total_incl_d", "TOTAL incl. module D (with benefits)"),
                     ("per_m2", "Intensity (kgCO₂e/m²)"),
                     ("per_m2_year", "Intensity (kgCO₂e/m²/yr)")]:
        ws.cell(row, 2, label).font = Font(bold=True)
        ws.cell(row, 3, r["totals"][k]).font = Font(bold=True)
        row += 1

    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 18

    # Materials sheet
    ms = wb.create_sheet("MATERIALS")
    for col, h in enumerate(["Material", "Mass (t)", "A1-A3 (tCO₂e)"], start=1):
        c = ms.cell(1, col, h); c.font = _HEAD; c.fill = _HFILL
    for i, m in enumerate(r["detail"], start=2):
        ms.cell(i, 1, m["name"]); ms.cell(i, 2, m["mass_t"]); ms.cell(i, 3, m["a1a3_t"])
    ms.column_dimensions["A"].width = 40
    ms.column_dimensions["B"].width = 12
    ms.column_dimensions["C"].width = 16
    ms["A1"].alignment = Alignment(horizontal="left")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
